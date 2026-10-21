"""Render the unchanged public Excel template as readable, local HTML sheets.

Reads the workbook only; never opens private trading workbooks or writes Excel.
Requires openpyxl for extraction. No broker or market-data calls are made.
"""
from pathlib import Path
from html import escape
from datetime import date, datetime
import hashlib, json, re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/files/US_Robust_Ops_Template_v1.xlsx'
META = ROOT / 'verification/ledger-preview.json'
meta = json.loads(META.read_text(encoding='utf-8'))
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == meta['sha256']
book = load_workbook(SOURCE, data_only=False)
assert book.sheetnames == meta['sheets']
assert not book._external_links
labels = ['설정·기록', '보유 현황', '종목 선정', '성과', '점검', '주문']

def color(value, fallback):
    if value and value.type == 'rgb' and re.fullmatch(r'[A-Fa-f0-9]{8}', value.rgb):
        return '#' + value.rgb[-6:]
    return fallback

def display(cell):
    value = cell.value
    if value is None: return ''
    if isinstance(value, (date, datetime)): return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        fmt = cell.number_format
        digits = len(re.search(r'\.([0#]+)', fmt).group(1)) if re.search(r'\.([0#]+)', fmt) else 0
        if '%' in fmt: return format(value, f'.{digits}%')
        if fmt != 'General': return format(value, f',.{digits}f' if ',' in fmt else f'.{digits}f')
    return str(value)

tabs, sections, counts = [], [], {}
for i, sheet in enumerate(book):
    tabs.append(f'<button type="button" role="tab" id="tab-{i}" aria-controls="sheet-{i}" aria-selected="{str(i == 0).lower()}" tabindex="{0 if i == 0 else -1}">{i:02d} {labels[i]}</button>')
    hidden, spans = set(), {}
    for merge in sheet.merged_cells.ranges:
        spans[(merge.min_row, merge.min_col)] = (merge.max_row-merge.min_row+1, merge.max_col-merge.min_col+1)
        hidden.update((r,c) for r in range(merge.min_row,merge.max_row+1) for c in range(merge.min_col,merge.max_col+1) if (r,c)!=(merge.min_row,merge.min_col))
    widths = [max(85, sheet.column_dimensions[get_column_letter(c)].width * 7) for c in range(1,sheet.max_column+1)]
    cols = '<col style="width:42px">' + ''.join(f'<col style="width:{width:.0f}px">' for width in widths)
    rows = ['<thead><tr><th></th>' + ''.join(f'<th scope="col">{get_column_letter(c)}</th>' for c in range(1,sheet.max_column+1)) + '</tr></thead><tbody>']
    counts[sheet.title] = 0
    for row in sheet:
        cells = [f'<th scope="row">{row[0].row}</th>']
        for cell in row:
            if (cell.row,cell.column) in hidden: continue
            if cell.value is not None: counts[sheet.title] += 1
            rs, cs = spans.get((cell.row,cell.column),(1,1))
            bg = color(cell.fill.fgColor,'#fff') if cell.fill.patternType == 'solid' else '#fff'
            fg = color(cell.font.color,'#202d3d')
            align = cell.alignment.horizontal
            if align not in ('left','center','right'): align = 'right' if isinstance(cell.value,(int,float)) else 'left'
            style = f'background:{bg};color:{fg};font-weight:{700 if cell.font.bold else 400};text-align:{align}'
            raw = json.dumps(cell.value,ensure_ascii=False,default=str)
            cells.append(f'<td data-cell="{cell.coordinate}" data-raw="{escape(raw,quote=True)}" rowspan="{rs}" colspan="{cs}" style="{style}">{escape(display(cell))}</td>')
        rows.append('<tr>'+''.join(cells)+'</tr>')
    rows.append('</tbody>')
    sections.append(f'<section id="sheet-{i}" role="tabpanel" aria-labelledby="tab-{i}" {"hidden" if i else ""}><h2>{escape(sheet.title)}</h2><div class="grid" tabindex="0" role="region" aria-label="{escape(sheet.title)} 시트, 가로 스크롤 가능"><table style="width:{42+sum(widths):.0f}px"><caption>{escape(sheet.title)} · 공개 템플릿</caption><colgroup>{cols}</colgroup>{"".join(rows)}</table></div></section>')
assert counts == meta['nonempty_cells']
html = '''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Excel 원장 샘플 | 김재원 포트폴리오</title><style>
*{box-sizing:border-box}body{margin:0;color:#202d3d;background:#fff;font-family:"Malgun Gothic",sans-serif;word-break:keep-all}main{max-width:1500px;margin:auto;padding:32px 28px}a{color:#006f6b;text-underline-offset:4px}header{border-bottom:2px solid #006f6b;padding-bottom:20px}h1{color:#006f6b;font-size:27px;margin:20px 0 12px}p{line-height:1.8;font-size:13px;margin:8px 0}.notice{background:#f5f7f9;border-left:3px solid #006f6b;padding:14px 18px;margin:20px 0}.actions{display:flex;flex-wrap:wrap;gap:18px;font-size:13px}.tabs{display:flex;gap:7px;overflow-x:auto;margin:26px 0 18px;padding:4px 2px 8px}button{font:inherit;white-space:nowrap;border:1px solid #dce3ea;background:white;padding:10px 16px;cursor:pointer;font-size:13px;color:#526477}button[aria-selected=true]{color:#fff;background:#006f6b;border-color:#006f6b}button:focus-visible,a:focus-visible,.grid:focus-visible{outline:2px solid #007d76;outline-offset:3px}h2{font-size:16px;margin:0 0 12px}.grid{overflow:auto;border:1px solid #dce3ea;max-height:70vh}table{table-layout:fixed;border-collapse:separate;border-spacing:0;min-width:100%;font-size:12px;font-variant-numeric:tabular-nums}caption{position:absolute;width:1px;height:1px;overflow:hidden;clip-path:inset(50%)}th,td{padding:8px 9px;border-right:1px solid #e0e5eb;border-bottom:1px solid #e0e5eb;min-height:30px;line-height:1.6;overflow-wrap:anywhere;white-space:pre-wrap;vertical-align:middle}th{background:#f1f4f7;font-weight:400;color:#687481;text-align:center}thead th{position:sticky;top:0;z-index:2}tbody th{position:sticky;left:0;z-index:1}footer{font-size:12px;color:#687481;line-height:1.8;margin-top:20px}[hidden]{display:none!important}@media(max-width:600px){main{padding:22px 16px}h1{font-size:23px}.notice{padding:12px 14px}button{padding:9px 12px}.grid{max-height:65vh}}
</style></head><body><main><header><a href="../index.html#operations">← 포트폴리오로 돌아가기</a><h1>Excel 원장 샘플</h1><p>공개 Excel 템플릿의 6개 시트를 브라우저에서 확인할 수 있습니다.</p><div class="actions"><a href="../files/US_Robust_Ops_Template_v1.xlsx" download>Excel 원본 다운로드</a><a href="SOURCE_URL" target="_blank" rel="noopener">공개 원본 출처</a></div></header><div class="notice"><p><strong>공개 템플릿에 포함된 예시 값입니다.</strong> 실계좌 거래내역이나 실제 체결·수익률 증빙이 아닙니다.</p><p>포트폴리오의 7개 시트는 현재 생성 코드가 주간 추이 시트(06_Weekly_Trend)를 추가하는 구성입니다. 이 공개 템플릿에는 00~05의 6개 시트가 있습니다.</p></div><div class="tabs" role="tablist" aria-label="Excel 시트">TABS</div>SECTIONS<footer>시트 안에서 가로·세로로 스크롤할 수 있습니다. 셀 값은 원본 그대로이며, 계산·주문 실행 기능은 제공하지 않습니다. 원본 이용허락: <a href="../files/US_Robust_Ops_Template_LICENSE.txt">MIT License</a>.</footer></main><script>
const tabs=[...document.querySelectorAll('[role=tab]')];function select(i){tabs.forEach((t,n)=>{const active=n===i;t.setAttribute('aria-selected',String(active));t.tabIndex=active?0:-1;document.getElementById(t.getAttribute('aria-controls')).hidden=!active;});}tabs.forEach((t,i)=>{t.addEventListener('click',()=>select(i));t.addEventListener('keydown',e=>{let j;if(e.key==='ArrowRight')j=(i+1)%tabs.length;else if(e.key==='ArrowLeft')j=(i-1+tabs.length)%tabs.length;else if(e.key==='Home')j=0;else if(e.key==='End')j=tabs.length-1;else return;e.preventDefault();select(j);tabs[j].focus();});});
</script></body></html>'''.replace('SOURCE_URL',escape(meta['source'],quote=True)).replace('TABS',''.join(tabs)).replace('SECTIONS',''.join(sections))
target = ROOT/'docs/ledger/index.html'
target.parent.mkdir(parents=True,exist_ok=True)
target.write_text(html,encoding='utf-8')
print('Rendered public template:',len(book.sheetnames),'sheets,',sum(counts.values()),'nonempty cells; workbook bytes unchanged.')
